"""Regenerate the sample documents in tests/data.

The documents are committed; tests only read them. Run this only to change
them:  uv run python tests/data/make_test_documents.py
"""
from pathlib import Path
import struct

from docx import Document as DocxDocument

DATA_DIR = Path(__file__).resolve().parent

# Some download pages inject HTML before %PDF; the CD player manual had exactly
# this prefix, which made PyPDF2 crash with KeyError: '/Root'.
JUNK_PREFIX = b"<!==  template/download.html ==>"


def build_pdf(pages: list[str], xref_stream: bool = False) -> bytes:
    """Build a minimal valid PDF with one line of Helvetica text per page.

    `xref_stream=True` writes the cross-reference index as a stream object
    (PDF 1.5+, used by most modern PDFs including the CD player manual)
    instead of a plain `xref` table.
    """
    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] "
           f"/Count {len(pages)} >>".encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for text, page_id in zip(pages, page_ids):
        content_id = page_id + 1
        stream = f"BT /F1 12 Tf 72 720 Td ({escape(text)}) Tj ET".encode()
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        objects[content_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.5\n")
    offsets = {}
    for obj_id in sorted(objects):
        offsets[obj_id] = len(out)
        out += f"{obj_id} 0 obj\n".encode() + objects[obj_id] + b"\nendobj\n"
    xref_offset = len(out)

    if xref_stream:
        # Uncompressed xref stream: per object 1 byte type, 4 bytes offset, 2 bytes generation.
        xref_id = max(objects) + 1
        size = xref_id + 1
        offsets[xref_id] = xref_offset
        rows = b"\x00" + struct.pack(">IH", 0, 65535) + b"".join(
            b"\x01" + struct.pack(">IH", offsets[i], 0) for i in range(1, size))
        out += (f"{xref_id} 0 obj\n<< /Type /XRef /Size {size} /W [1 4 2] /Root 1 0 R "
                f"/Length {len(rows)} >>\nstream\n").encode() + rows + b"\nendstream\nendobj\n"
        out += f"startxref\n{xref_offset}\n%%EOF\n".encode()
        return bytes(out)

    size = max(objects) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for obj_id in range(1, size):
        out += f"{offsets[obj_id]:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {size} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n").encode()
    return bytes(out)


def guide_text() -> str:
    """A multi-section guide long enough to split into several 800-token chunks."""
    return "\n\n".join(
        f"Topic {i}: instructions for feature {i}, step {i * 2}, setting {i * 5}. " * 30
        for i in range(12)
    ) + "\n"


def main() -> None:
    (DATA_DIR / "notes.txt").write_text("Hello from a text file.\nSecond line.", encoding="utf-8")
    (DATA_DIR / "guide.txt").write_text(guide_text(), encoding="utf-8")
    (DATA_DIR / "data.csv").write_text("a,b\n1,2\n", encoding="utf-8")

    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    doc.save(str(DATA_DIR / "report.docx"))

    pdfs = {
        "manual.pdf": build_pdf(["The region number of this player is 2."]),
        "two-pages.pdf": build_pdf(["End of page one", "Start of page two"]),
        "xref-stream.pdf": build_pdf(["Modern PDF."], xref_stream=True),
        "junk-prefix-xref-table.pdf": JUNK_PREFIX + build_pdf(["Still readable."]),
        "junk-prefix-xref-stream.pdf": JUNK_PREFIX + build_pdf(["Still readable."], xref_stream=True),
    }
    for name, data in pdfs.items():
        (DATA_DIR / name).write_bytes(data)


if __name__ == "__main__":
    main()
